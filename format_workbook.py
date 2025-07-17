#!/usr/bin/python3
#
# This file is part of the GPM phenotyping scripts.
#
# Copyright (c) 2025 Jason Toney
#
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU General Public License as published by
# the Free Software Foundation, either version 3 of the License, or
# (at your option) any later version.
#
# This program is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
# GNU General Public License for more details.
#
# You should have received a copy of the GNU General Public License
# along with this program. If not, see <https://www.gnu.org/licenses/>.

"""docstring goes here."""


import os
import openpyxl


def format_workbook(workbook_file):
    """docstring goes here."""
    if os.path.exists(workbook_file):
        workbook = openpyxl.load_workbook(workbook_file)
    else:
        return
    # Auto-size columns to fit content
    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        for column in sheet.columns:
            max_length = 0
            column_letter = openpyxl.utils.get_column_letter(column[0].column)
            for cell in column:
                if cell.value == "Quality Check":
                    max_length = 13
                    break
                if len(str(cell.value)) > max_length:
                    max_length = len(cell.value)
            adjusted_width = max_length + 2
            sheet.column_dimensions[column_letter].width = adjusted_width
    workbook.save(workbook_file)


def main():
    """docstring goes here."""
    os.chdir(os.path.dirname(__file__))
    format_workbook("GPMPhenotypingAssay_Workbook.xlsx")


if __name__ == "__main__":
    main()

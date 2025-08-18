#!/usr/bin/env python3
#
# This file is part of the GPM phenotyping scripts.
#
# Copyright (c) 2025 Jason Toney
#
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU General Public License as published by
# the Free Software Foundation, either version 3 of the License, or
# (at your option) any later version.
#
# This program is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
# GNU General Public License for more details.
#
# You should have received a copy of the GNU General Public License
# along with this program. If not, see <https://www.gnu.org/licenses/>.

"""docstring goes here"""


import os
import re
import sys
import warnings

import matplotlib.pyplot as plt
import numpy as np
import openpyxl
import pandas as pd
from scipy.optimize import curve_fit
from sklearn.metrics import r2_score


def logistic_4pl(x, bottom, top, ed50, hill_slope):
    """docstring goes here"""
    return bottom + (top - bottom) / (
        1 + np.exp(hill_slope * (np.log(x) - np.log(ed50)))
    )


def extract_timepoint(plate):
    """docstring goes here"""
    matches = re.compile(r"(\d+(?:\.\d+)?)(?i:hr)").search(plate)
    return float(matches.group(1)) if matches else 1.0


def plot_curve(index, plate, concentrations, germination_rates, popt, pcov, color=True):
    """docstring goes here"""
    # Calculate ED50
    ed50 = int(round(popt[2], 0))
    # Calculate standard error for ED50
    ed50_SE = int(round(np.sqrt(np.diag(pcov))[2], 0))
    # Calculate R-squared
    fitted = logistic_4pl(concentrations, *popt)
    r2 = round(r2_score(germination_rates, fitted), 5)
    # Plot the DRC
    x_vals = np.linspace(min(concentrations), max(concentrations), 100)
    y_vals = logistic_4pl(x_vals, *popt)
    if index <= 0:
        plt.figure(figsize=(10, 6))
    colors = ["green", "blue", "red", "orange"]
    line_styles = ["-", "--", "-.", ":"]
    markers = ["o", "^", "s", "X"]
    # Add a separator to the legend for readability
    if index > 0:
        plt.plot([], [], "", label="-" * 30, linestyle="None", marker="")
    plt.scatter(
        concentrations,
        germination_rates,
        label="Data"
        if index < 0
        else f"{extract_timepoint(plate)}hr ({plate.split('-')[0]} - {plate.split('-')[1]})",
        color=colors[index] if (index >= 0 and color) else "black",
        marker=markers[index if index >= 0 else 0],
        edgecolors="black",
    )
    plt.plot(
        x_vals,
        y_vals,
        label="Fitted Curve",
        color=colors[index] if (index >= 0 and color) else "black",
        linestyle=line_styles[index if (index >= 0 and not color) else 0],
    )
    plt.axvline(
        ed50,
        linestyle=line_styles[index if (index >= 0 and not color) else 1],
        color=colors[index] if (index >= 0 and color) else "black",
        label=f"ED$_{{50}}$ = {ed50} ± {ed50_SE}",
    )
    # add R-squared value to the legend
    plt.plot([], [], "", label=f"R$^2$ = {r2}", linestyle="None", marker="")
    min_tick = int(np.floor(min(np.min(concentrations), 0) / 50.0) * 50)
    max_tick = int(np.ceil(max(np.max(concentrations), ed50) / 50.0) * 50)
    plt.xticks(np.arange(min_tick, max_tick + 1, 50))
    plt.xlabel("UV-C Dose (J/m$^2$)")
    plt.ylabel("Mean germination relative to control (%)")
    plt.legend()
    return ed50, ed50_SE


def calculate_ed50(isolates, plates, show_plot=False):
    """docstring goes here"""
    # ignore annoying warnings
    warnings.filterwarnings("ignore")
    # Plot up to 4 runs at a time
    if len(isolates) > 4:
        print(f"Can only plot up to 4 runs at a time, {len(isolates)} runs selected.")
        return
    # Load the data
    workbook = "GPMPhenotypingAssay_Workbook.xlsx"
    if os.path.exists(workbook):
        uvc_workbook = openpyxl.load_workbook(workbook)
    else:
        print(f"Missing workbook: {workbook}")
        return
    for index, (isolate, plate) in enumerate(zip(isolates, plates)):
        # Format isolate name and plate ID
        isolate = isolate.upper()
        plate = plate.upper()
        sheet_name = f"{isolate} ({plate})"
        if sheet_name in uvc_workbook.sheetnames:
            sheet = uvc_workbook[sheet_name]
        else:
            print(f"Missing sheet: {isolate} {plate}")
            return
        # Load the data
        data = pd.DataFrame(sheet.values)
        data.columns = data.iloc[0]
        # Take the average of the controls
        controls = data[data["Treatment"].str.contains("Control")]
        control_avg = sum(controls["48hr %"].values) / len(controls)
        # Extract dose and response data
        data = data[data["Treatment"].str.contains("Speed|J/m2", na=False)]
        data["Concentration"] = data["Treatment"].str.extract(r"(\d+)").astype(int)
        concentrations = data["Concentration"].values
        germination_rates = data["48hr %"].values
        # Normalize germination rates relative to the controls
        for i, n in enumerate(germination_rates):
            germination_rates[i] = min(round(n / control_avg * 100, 2), 100)
        # Use mean germination values
        size = int(len(germination_rates) / len(set(concentrations)))
        concentrations = [
            concentrations[i] for i in range(0, len(concentrations), size)
        ]
        germination_rates = [
            sum(germination_rates[i : i + size]) / size
            for i in range(0, len(germination_rates), size)
        ]
        # Fit the curve and generate ED50
        initial_guess = [
            min(germination_rates),
            max(germination_rates),
            np.median(concentrations),
            -1,
        ]
        bounds = (
            [0, 0, min(concentrations), -10],
            [100, 100, max(concentrations) * 10, 10],
        )
        popt, pcov = curve_fit(  # pylint: disable=unbalanced-tuple-unpacking
            logistic_4pl,
            concentrations,
            germination_rates,
            p0=initial_guess,
            maxfev=10000,
            bounds=bounds,
        )
        # Handle plotting of a single assay run
        if len(isolates) == 1:
            index = -1
        ed50, ed50_SE = plot_curve(
            index, plate, concentrations, germination_rates, popt, pcov
        )
        print(f"Estimated UV-C ED50: {isolate}, {plate}: {ed50} ± {ed50_SE} J/m^2")
        # add the ED50 to tracking spreadsheet
        if "Assay Data" in uvc_workbook.sheetnames:
            sheet = uvc_workbook["Assay Data"]
        else:
            print("Missing sheet: Assay Data")
            return
        assay_df = pd.DataFrame(sheet.values)
        assay_df.columns = assay_df.iloc[0]
        for idx, row in assay_df[1:].iterrows():
            if row["Isolate"] == isolate and row["Plate ID"].upper() == plate:
                sheet.cell(
                    row=idx + 1,
                    column=assay_df.columns.get_loc("ED50 (J/m^2)") + 1,
                    value=f"{ed50} ± {ed50_SE}",
                )
        uvc_workbook.save(workbook)
    if len(plates) == 1:
        plt.title(f"UV-C Dose-Response Curve: {isolates[0]} - {plates[0]}")
        plt.savefig(
            f"results/ED50_{isolates[0]}_{plates[0]}.png", dpi=300, bbox_inches="tight"
        )
    else:
        plotted = (
            " - ".join(i for i in isolates) if len(set(isolates)) > 1 else isolates[0]
        )
        plt.title(f"UV-C Dose-Response Curve: {plotted}")
        plt.savefig(
            f"results/ED50_{plotted}_combined.png", dpi=300, bbox_inches="tight"
        )
    # Show the plot if requested
    if show_plot:
        plt.show()
    plt.close()


def main(isolates, plates, show_plot=False):
    """docstring goes here"""
    # Only calculate ED50 for UV-C assay runs
    for plate in plates:
        if "UVC" not in plate.upper():
            print(f"NOT A UVC PLATE: {plate}")
            return
    # Make sure equal number of isolates/plates specified
    if len(isolates) != len(plates):
        return
    os.chdir(os.path.dirname(__file__))
    # Sort plate IDs by timepoint while maintaining isolate-plate pairs
    paired = sorted(list(zip(plates, isolates)), key=lambda x: extract_timepoint(x[0]))
    plates, isolates = zip(*paired)
    plates, isolates = list(plates), list(isolates)
    calculate_ed50(isolates, plates, show_plot)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])

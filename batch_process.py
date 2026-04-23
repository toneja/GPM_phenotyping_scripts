#!/usr/bin/env python3
#
# This file is part of the GPM phenotyping scripts.
#
# Copyright (c) 2025 Jason Toney
#
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU General Public License as published by
# the Free Software Foundation, either version 3 of the License, or
# (at your option) any later version.
#
# This program is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
# GNU General Public License for more details.
#
# You should have received a copy of the GNU General Public License
# along with this program. If not, see <https://www.gnu.org/licenses/>.

"""
    This script executes the ImageJ macro and python scripts in a batch process.
    Input: all images found within the local "ECHO Images" folder.
    Output: XLSX file containing the total counts of conidia germination states.
"""


import os
import platform
import subprocess
import sys
import time
from datetime import datetime

import openpyxl
import pandas as pd
from PIL import Image

import analyze_results
import calculate_deltas
import calculate_ed50
import check_results
import compile_workbook
import convert_timings
import format_workbook


def batch_process(image_folder="ECHO Images", prompt=True):
    """Analyze all the images found in "ECHO Images" subdirectories."""
    # Start the timer
    start_time = time.time()
    # Count how many albums are processed
    processed = 0
    # Work inside this directory
    os.chdir(os.path.dirname(__file__))
    # Iterate through the image folders
    for folder_name in os.listdir(image_folder):
        # Full path to the current image folder
        current_folder = os.path.join(image_folder, folder_name)

        # Check if the current item is a directory
        if os.path.isdir(current_folder) and "plate" in current_folder:
            # Check if the album has already been processed
            if os.path.exists(f"ImageJ/GPM/images/{folder_name}") and os.path.exists(
                f"ImageJ/GPM/results/{folder_name}"
            ):
                print(f"Skipping folder: {current_folder}, already processed.")
                continue

            # make output folders
            os.makedirs(f"ImageJ/GPM/images/{folder_name}", exist_ok=True)
            os.makedirs(f"ImageJ/GPM/results/{folder_name}", exist_ok=True)

            print(f"Processing folder: {current_folder}")
            if any(x in folder_name for x in ("48hr", "Control")):
                processed += 1

            # Convert any tif files to jpg; req for ImageJ
            for file in os.listdir(current_folder):
                file = os.path.join(current_folder, file)
                if file.endswith(".tif"):
                    with Image.open(file) as tif_file:
                        tif_file.convert("RGB").save(
                            file.replace(".tif", ".jpg"), "JPEG"
                        )

            # Execute the ImageJ macro for the current folder
            imagej_bin = (
                "/Applications/Fiji.app/Contents/MacOS/ImageJ-macosx"
                if platform.system() == "Darwin"
                else "./ImageJ/ImageJ.exe"
            )
            command = [
                imagej_bin,
                "-macro",
                "ImageJ/GPM/BatchProcess.ijm",
                current_folder,
            ]

            try:
                subprocess.run(command, capture_output=True, text=True, check=True)
            except subprocess.CalledProcessError as exception:
                print(f"Error executing the macro: {exception}")

            # Relocate output files into their respective release folders
            for file in os.listdir("ImageJ/GPM/images"):
                if file.endswith(".tif"):
                    os.replace(
                        f"ImageJ/GPM/images/{file}",
                        f"ImageJ/GPM/images/{folder_name}/{file}",
                    )
            for file in os.listdir("ImageJ/GPM/results"):
                if file.endswith(".csv"):
                    os.replace(
                        f"ImageJ/GPM/results/{file}",
                        f"ImageJ/GPM/results/{folder_name}/{file}",
                    )

    # Process the ImageJ results
    for folder in os.listdir("ImageJ/GPM/results"):
        if folder.endswith("48hr"):
            analyze_results.main(f"ImageJ/GPM/results/{folder}")

    # Check the results for insufficient germination and spore deposition
    check_results.main()

    # Convert manual distance timings to real-world exposure output units
    convert_timings.main()

    # Compile the results into a workbook
    compile_workbook.main()

    # Generate Dose-Response curves and update workbook info
    workbook = "GPMPhenotypingAssay_Workbook.xlsx"
    if os.path.exists(workbook):
        uvc_workbook = openpyxl.load_workbook(workbook)
        sheet_name = "Assay Data"
        if sheet_name in uvc_workbook.sheetnames:
            sheet = uvc_workbook[sheet_name]
            df = pd.DataFrame(sheet.values)
            df.columns = df.iloc[0]
            for idx, row in df[1:].iterrows():
                plate = row["Plate ID"]
                # Only calculate ED50 if it is missing from the workbook
                if row["Quality Check"] == "PASS" and pd.isna(row["ED50 (J/m^2)"]):
                    calculate_ed50.main([row["Isolate"]], [plate], False)
                # Fill in assay data from plate ID string
                date = datetime.strptime(plate.split("-")[1], "%m%d%y").strftime(
                    "%m/%d/%Y"
                )
                timepoint = plate.split("-")[2].split("hr")[0]
                timepoint += " hour" if float(timepoint) == 1 else " hours"
                wavelength = plate.split("-")[3]
                sheet.cell(
                    row=idx + 1, column=df.columns.get_loc("Date") + 1, value=date
                )
                sheet.cell(
                    row=idx + 1,
                    column=df.columns.get_loc("App Time") + 1,
                    value=timepoint,
                )
                sheet.cell(
                    row=idx + 1,
                    column=df.columns.get_loc("Wavelength") + 1,
                    value=wavelength,
                )
            uvc_workbook.save(workbook)

    # Calculate Deltas for new UV-C runs
    calculate_deltas.main()

    # Fix up the workbook formatting
    format_workbook.main()

    # Calculate the elapsed time
    elapsed_time = time.time() - start_time
    # Print elapsed time in H:M:S format
    print(f"\nElapsed time: {time.strftime('%H:%M:%S', time.gmtime(elapsed_time))}")
    print(f"Assay runs processed: {processed}")
    if prompt:
        input("Batch processing complete. Press ENTER.\n")


if __name__ == "__main__":
    if len(sys.argv) > 1:
        IMAGE_FOLDER = sys.argv[1]
    else:
        IMAGE_FOLDER = "ECHO Images"
    batch_process(IMAGE_FOLDER, True)
